import os
from datetime import datetime, date, time
import openpyxl
from app.db.session import SessionLocal, db_manager, Base
import app.models.entities
from app.models.entities import Incidente, Aplicacion

def seed_incidentes_from_excel(excel_path: str = r"d:\Proyectos_Anti\Detalle_JOBS.xlsx", db=None):
    close_db = False
    if db is None:
        db = SessionLocal()
        close_db = True

    try:
        # Asegurar tablas creadas
        Base.metadata.create_all(bind=db_manager.engine)

        if not os.path.exists(excel_path):
            print(f"Archivo Excel de incidentes no encontrado en {excel_path}")
            return

        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet = wb["JOB_DTSX"] if "JOB_DTSX" in wb.sheetnames else wb.active

        # 1. Asegurar aplicación FCD en catálogo APLICACIONES si no existe
        app_exist = db.query(Aplicacion).filter(Aplicacion.SIGLAS.ilike("FCD")).first()
        if not app_exist:
            app_nueva = Aplicacion(
                NOMBRE_APLICACION="FINANCIAMIENTO DE VENTAS",
                SIGLAS="FCD",
                ESTADO="Activo",
                FECHA_REGISTRO=datetime.now()
            )
            db.add(app_nueva)
            db.commit()

        # 2. Importar incidentes si la tabla está vacía
        if db.query(Incidente).count() == 0:
            print("Importando registros iniciales de Detalle_JOBS.xlsx (JOB_DTSX)...")
            insert_count = 0
            for r in range(2, sheet.max_row + 1):
                vals = [sheet.cell(row=r, column=c).value for c in range(1, 14)]
                if not any(v is not None for v in vals):
                    continue

                atendido = str(vals[0]).strip() if vals[0] else None
                app_val = str(vals[1]).strip() if vals[1] else "FCD"

                ruta_critica = str(vals[2]).strip().upper() if vals[2] else "NO"
                if "SI" in ruta_critica:
                    ruta_critica = "SI"
                elif "NO" in ruta_critica:
                    ruta_critica = "NO"
                else:
                    ruta_critica = "NO"

                job_val = str(vals[3]).strip() if vals[3] is not None else f"JOB-{r}"

                f_canc = vals[4]
                if isinstance(f_canc, datetime):
                    f_canc = f_canc.date()
                elif not isinstance(f_canc, date):
                    f_canc = None

                server_val = str(vals[5]).strip() if vals[5] else None
                ruta_val = str(vals[6]).strip() if vals[6] else None
                dtsx_val = str(vals[7]).strip() if vals[7] else None
                aplicar_val = str(vals[8]).strip() if vals[8] else None

                h_canc = vals[9]
                if isinstance(h_canc, time):
                    h_canc_str = h_canc.strftime("%H:%M")
                elif h_canc is not None:
                    h_canc_str = str(h_canc).strip()
                else:
                    h_canc_str = None

                desc_error = str(vals[10]).strip() if vals[10] else None
                solucion_val = str(vals[11]).strip() if vals[11] else None

                f_sol = vals[12]
                if isinstance(f_sol, datetime):
                    pass
                elif isinstance(f_sol, date):
                    f_sol = datetime.combine(f_sol, datetime.min.time())
                else:
                    f_sol = None

                ambiente_val = "PRD"

                nuevo_inc = Incidente(
                    ATENDIDO_POR=atendido,
                    APLICATIVO=app_val,
                    RUTA_CRITICA=ruta_critica,
                    JOB=job_val,
                    FECHA_CANCELACION=f_canc,
                    SERVER=server_val,
                    RUTA=ruta_val,
                    DTSX=dtsx_val,
                    APLICAR=aplicar_val,
                    HORA_CANCELACION=h_canc_str,
                    DESCRIPCION_ERROR=desc_error,
                    SOLUCION=solucion_val,
                    FECHA_HORA_SOLUCION=f_sol,
                    AMBIENTE=ambiente_val,
                    FECHA_CREACION=datetime.now(),
                    FECHA_ACTUALIZACION=datetime.now()
                )
                db.add(nuevo_inc)
                insert_count += 1

            db.commit()
            print(f"Se importaron {insert_count} incidentes desde {excel_path} exitosamente.")

    except Exception as e:
        print(f"Error en seed_incidentes_from_excel: {e}")
        db.rollback()
    finally:
        if close_db:
            db.close()

if __name__ == "__main__":
    seed_incidentes_from_excel()
