import os
from datetime import datetime, date
import openpyxl
from app.db.session import SessionLocal, db_manager, Base
import app.models.entities
from app.models.entities import Pase, Proyecto, Aplicacion

def seed_pases_from_excel(excel_path: str = r"d:\Proyectos_Anti\PASES_IBK.xlsx", db=None):
    close_db = False
    if db is None:
        db = SessionLocal()
        close_db = True

    try:
        # Asegurar tablas creadas
        Base.metadata.create_all(bind=db_manager.engine)

        if not os.path.exists(excel_path):
            print(f"Archivo Excel no encontrado en {excel_path}")
            return

        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet = wb["PASES2026"] if "PASES2026" in wb.sheetnames else wb.active

        # Mapeo de estados de Excel a los permitidos:
        # REGISTRADO, UAT_SOLICITADO, UAT_DESPLEGADO, QA_CERTIFICADO, PRD_SOLICITADO, PRD_EJECUTADO, RECHAZADO, ANULADO
        state_map = {
            'PRD_DESPLEGADO': 'PRD_EJECUTADO',
            'PRD_Desplegando': 'PRD_SOLICITADO',
            'PRD_EJECUTADO': 'PRD_EJECUTADO',
            'PRD_SOLICITADO': 'PRD_SOLICITADO',
            'UAT_DESPLEGADO': 'UAT_DESPLEGADO',
            'uat_desplegado': 'UAT_DESPLEGADO',
            'UAT_Desplegado': 'UAT_DESPLEGADO',
            'UAT_Solicitado': 'UAT_SOLICITADO',
            'UAT_SOLICITADO': 'UAT_SOLICITADO',
            'QA_Certificado': 'QA_CERTIFICADO',
            'QA_CERTIFICADO': 'QA_CERTIFICADO',
            'Registrado': 'REGISTRADO',
            'REGISTRADO': 'REGISTRADO',
            'RECHAZADO': 'RECHAZADO',
            'ANULADO': 'ANULADO'
        }

        # 1. Asegurar proyectos del Excel en catálogo PROYECTOS si no existen
        proyectos_excel = ["LPC", "WBC", "BSE", "FCD", "PUN-FCD", "C2C-FCD"]
        for p_nom in proyectos_excel:
            p_exist = db.query(Proyecto).filter(Proyecto.NOMBRE_PROYECTO.ilike(p_nom)).first()
            if not p_exist:
                p_nuevo = Proyecto(
                    NOMBRE_PROYECTO=p_nom,
                    DESCRIPCION_PROYECTO=f"Proyecto {p_nom} importado de catálogo de pases",
                    EQUIPO_SOLICITANTE="Sistemas / TI",
                    ESTADO="Activo",
                    FECHA_REGISTRO=datetime.now()
                )
                db.add(p_nuevo)
                db.commit()

        # 2. Asegurar aplicaciones del Excel en catálogo APLICACIONES si no existen
        apps_excel = [
            ("LPC", "LINEA PARALELA DE CREDITO"),
            ("WBC", "WEB BANKING CORPORATIVO"),
            ("FCD", "FINANCIAMIENTO DE VENTAS"),
            ("BSE", "BANCA SERVICIOS ELECTRONICOS")
        ]
        for siglas, nombre in apps_excel:
            app_exist = db.query(Aplicacion).filter(Aplicacion.SIGLAS.ilike(siglas)).first()
            if not app_exist:
                app_nueva = Aplicacion(
                    NOMBRE_APLICACION=nombre,
                    SIGLAS=siglas,
                    ESTADO="Activo",
                    FECHA_REGISTRO=datetime.now()
                )
                db.add(app_nueva)
                db.commit()

        # 3. Importar pases si la tabla está vacía
        if db.query(Pase).count() == 0:
            print("Importando registros de PASES2026...")
            insert_count = 0
            for r in range(2, sheet.max_row + 1):
                vals = [sheet.cell(row=r, column=c).value for c in range(1, 23)]
                if not any(v is not None for v in vals):
                    continue

                tipo_amb = str(vals[0]).strip().upper() if vals[0] else 'D'
                if tipo_amb not in ('A', 'D', 'H'):
                    tipo_amb = 'D'
                proy = str(vals[1]).strip() if vals[1] else None
                app_val = str(vals[2]).strip() if vals[2] else None

                f_reg_srt = vals[3].date() if isinstance(vals[3], datetime) else (vals[3] if isinstance(vals[3], date) else None)
                f_sol_uat = vals[4].date() if isinstance(vals[4], datetime) else (vals[4] if isinstance(vals[4], date) else None)
                f_des_uat = vals[5].date() if isinstance(vals[5], datetime) else (vals[5] if isinstance(vals[5], date) else None)

                st_raw = str(vals[6]).strip() if vals[6] else 'REGISTRADO'
                st = state_map.get(st_raw, 'REGISTRADO')

                cod_srt = str(vals[7]).strip() if vals[7] else None
                titulo = str(vals[8]).strip() if vals[8] else '(Sin título)'
                conf_prd = str(vals[9]).strip() if vals[9] else None
                qa = str(vals[10]).strip() if vals[10] else None
                f_cert = vals[11].date() if isinstance(vals[11], datetime) else (vals[11] if isinstance(vals[11], date) else None)
                f_reg_oc = vals[12].date() if isinstance(vals[12], datetime) else (vals[12] if isinstance(vals[12], date) else None)

                f_h_prd = vals[13] if isinstance(vals[13], datetime) else None
                oc = str(vals[14]).strip() if vals[14] is not None else None
                st_oc = str(vals[15]).strip() if vals[15] is not None else None
                op_pase = str(vals[16]).strip() if vals[16] is not None else None
                dev = str(vals[17]).strip() if vals[17] is not None else None
                sustento = str(vals[18]).strip() if vals[18] is not None else None
                user_aprob = str(vals[19]).strip() if vals[19] is not None else None
                q_sp = str(vals[20]).strip() if vals[20] is not None else None
                owner = str(vals[21]).strip() if vals[21] is not None else None

                pase = Pase(
                    TIPO_AMBIENTE=tipo_amb,
                    PROYECTO=proy,
                    APP=app_val,
                    FECHA_REGISTRO_SRT=f_reg_srt,
                    FECHA_SOLICITADO_UAT=f_sol_uat,
                    FECHA_DESPLEGADO_UAT=f_des_uat,
                    ESTADO_SRT=st,
                    CODIGO_SRT=cod_srt,
                    TITULO=titulo,
                    CONFORMES_PRD=conf_prd,
                    QA=qa,
                    FECHA_CERTIFICACION=f_cert,
                    FECHA_REGISTRO_OC=f_reg_oc,
                    FECHA_HORA_PASE_PRD=f_h_prd,
                    OC=oc,
                    STADO_OC=st_oc,
                    OPERADOR_PASE=op_pase,
                    DEV=dev,
                    SUSTENTO_VALOR_NEGOCIO=sustento,
                    USUARIO_FINAL_APROBACION=user_aprob,
                    Q_SP_PRD=q_sp,
                    RESPONSABLE_OWNER_PROYECTO=owner,
                    MOTIVO_ESTADO=None,
                    FECHA_REGISTRO=datetime.now(),
                    FECHA_ACTUALIZACION=datetime.now()
                )
                db.add(pase)
                insert_count += 1

            db.commit()
            print(f"Se insertaron {insert_count} pases exitosamente desde el Excel.")
        else:
            print(f"La tabla PASES ya contiene {db.query(Pase).count()} registros. No se duplicaron.")

    finally:
        if close_db:
            db.close()

if __name__ == "__main__":
    seed_pases_from_excel()
