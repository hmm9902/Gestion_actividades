import os
from datetime import datetime, date
import openpyxl
from app.db.session import SessionLocal, db_manager, Base
import app.models.entities
from app.models.entities import Ticket, Proyecto, Aplicacion

def seed_tickets_from_excel(excel_path: str = r"d:\Proyectos_Anti\GESTION_TICKETS.xlsx", db=None):
    close_db = False
    if db is None:
        db = SessionLocal()
        close_db = True

    try:
        # Asegurar tablas creadas
        Base.metadata.create_all(bind=db_manager.engine)

        if not os.path.exists(excel_path):
            print(f"Archivo Excel de tickets no encontrado en {excel_path}")
            return

        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet = wb["Tickets"] if "Tickets" in wb.sheetnames else wb.active

        # 1. Asegurar proyectos del Excel en catálogo PROYECTOS si no existen
        proyectos_excel = ["C2C-FCD"]
        for p_nom in proyectos_excel:
            p_exist = db.query(Proyecto).filter(Proyecto.NOMBRE_PROYECTO.ilike(p_nom)).first()
            if not p_exist:
                p_nuevo = Proyecto(
                    NOMBRE_PROYECTO=p_nom,
                    DESCRIPCION_PROYECTO=f"Proyecto {p_nom} importado de catálogo de tickets",
                    EQUIPO_SOLICITANTE="Sistemas / TI",
                    ESTADO="Activo",
                    FECHA_REGISTRO=datetime.now()
                )
                db.add(p_nuevo)
                db.commit()

        # 2. Asegurar aplicaciones del Excel en catálogo APLICACIONES si no existen
        apps_excel = [
            ("FCD", "FINANCIAMIENTO DE VENTAS")
        ]
        for siglas, nombre in apps_excel:
            app_exist = db.query(Aplicacion).filter(Aplicacion.SIGLAS.ilike(siglas)).first()
            if not app_exist:
                app_nueva = Aplicacion(
                    NOMBRE_APLICACION=nombre,
                    SIGLAS=siglas,
                    ESTADO="Activo",
                    FECHA_REGISTRO=datetime.now()
                )
                db.add(app_nueva)
                db.commit()

        # 3. Mapeo de estados del Excel
        # Estados válidos: ABIERTO, ASIGNADO, ANULADO, RECHAZADO, DEVUELTO, EN_PROCESO, SOLUCIONADO
        def normalizar_estado(raw_val) -> str:
            if not raw_val:
                return "ABIERTO"
            s = str(raw_val).strip()
            s_upper = s.upper()
            if "CANCELADO" in s_upper or "ANULADO" in s_upper:
                return "ANULADO"
            if "CERRADO" in s_upper or "SOLUCIONADO" in s_upper:
                return "SOLUCIONADO"
            if "RECHAZADO" in s_upper:
                return "RECHAZADO"
            if "DEVUELTO" in s_upper:
                return "DEVUELTO"
            if "PROCESO" in s_upper:
                return "EN_PROCESO"
            if "ASIGNADO" in s_upper:
                return "ASIGNADO"
            return "ABIERTO"

        # 4. Importar tickets si la tabla está vacía
        if db.query(Ticket).count() == 0:
            print("Importando registros iniciales de GESTION_TICKETS.xlsx...")
            insert_count = 0
            for r in range(2, sheet.max_row + 1):
                vals = [sheet.cell(row=r, column=c).value for c in range(1, 13)]
                if not any(v is not None for v in vals):
                    continue

                tipo_val = str(vals[0]).strip() if vals[0] else 'Incident'
                if tipo_val not in ('Incident', 'Request', 'OC'):
                    tipo_val = 'Incident'

                f_reg = vals[1] if isinstance(vals[1], datetime) else (
                    datetime.combine(vals[1], datetime.min.time()) if isinstance(vals[1], date) else None
                )

                app_val = str(vals[2]).strip() if vals[2] else 'FCD'
                proy_val = str(vals[3]).strip() if vals[3] else 'C2C-FCD'
                
                amb_val = str(vals[4]).strip().upper() if vals[4] else None
                if amb_val and amb_val not in ('UAT', 'PRD'):
                    amb_val = 'UAT' if 'UAT' in amb_val else ('PRD' if 'PRD' in amb_val else None)

                ticket_code = str(vals[5]).strip() if vals[5] is not None else f"TKT-{r}"
                desc_val = str(vals[6]).strip() if vals[6] else None

                f_aten = vals[7] if isinstance(vals[7], datetime) else (
                    datetime.combine(vals[7], datetime.min.time()) if isinstance(vals[7], date) else None
                )

                ibm_asig = str(vals[8]).strip() if vals[8] else None
                cel_contacto = str(vals[9]).strip() if vals[9] else None
                estado_val = normalizar_estado(vals[10])
                comentario_val = str(vals[11]).strip() if vals[11] else None

                nuevo_tkt = Ticket(
                    TIPO=tipo_val,
                    FECHA_REGISTRO=f_reg,
                    APLICATIVO=app_val,
                    PROYECTO=proy_val,
                    AMBIENTE=amb_val,
                    TICKET=ticket_code,
                    DESCRIPCION=desc_val,
                    FECHA_ATENCION=f_aten,
                    IBM_ASIGNADO=ibm_asig,
                    CEL_CONTACTO=cel_contacto,
                    ESTADO=estado_val,
                    COMENTARIO=comentario_val,
                    FECHA_CREACION=datetime.now(),
                    FECHA_ACTUALIZACION=datetime.now()
                )
                db.add(nuevo_tkt)
                insert_count += 1

            db.commit()
            print(f"Se importaron {insert_count} tickets desde {excel_path} exitosamente.")

    except Exception as e:
        print(f"Error en seed_tickets_from_excel: {e}")
        db.rollback()
    finally:
        if close_db:
            db.close()

if __name__ == "__main__":
    seed_tickets_from_excel()
